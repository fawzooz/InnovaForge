"""The InnovaForge Toolkit as a fillable workbook, one per language.

    python3 project/toolkit/innovaforge/workbook.py

Reads out/content.json (written by build.mjs) and writes
out/innovaforge-toolkit-{en,ar}.xlsx: the canvas, the tabular sheets of every
stage, the five exit gates, with scores, risk ratings and progress computed.
Needs openpyxl.
"""
import json
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(DIR, 'out')
C = json.load(open(os.path.join(OUT, 'content.json'), encoding='utf-8'))

DARK, AMBER, AMBER_D, CREAM, LINE, PAPER = '111418', 'F5A623', 'A86A10', 'FBF3E3', 'E3DED4', 'F5F3EE'
thin = Side(style='thin', color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

L = 'en'


def t(o):
    return o if isinstance(o, str) else o[L]


def bi(en, ar):
    return {'en': en, 'ar': ar}


def num(x):
    s = str(x)
    return s.translate(str.maketrans('0123456789.', '٠١٢٣٤٥٦٧٨٩٫')) if L == 'ar' else s


class Sheet:
    """A worksheet written top to bottom, with the toolkit's styles."""

    def __init__(self, wb, code, title, instr, widths):
        name = f'{num(code)} {t(title)}' if code else t(title)
        self.ws = wb.create_sheet(name[:31])
        self.ws.sheet_view.rightToLeft = L == 'ar'
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_properties.tabColor = AMBER if code else DARK
        self.font = 'Noto Kufi Arabic' if L == 'ar' else 'Manrope'
        self.align = 'right' if L == 'ar' else 'left'
        self.ncols = len(widths)
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.r = 1
        self.text(t(C['meta']['runner']) + (f' · {num(code)}' if code else ''), size=9, color='8E8B85')
        self.text(t(title), size=16, bold=True, color=DARK, height=28)
        if instr:
            self.text(t(instr), size=10, color='424244', height=32, wrap=True)
        self.r += 1

    def f(self, size=10, bold=False, color='23272C'):
        return Font(name=self.font, size=size, bold=bold, color=color)

    def text(self, value, size=10, bold=False, color='23272C', height=None, wrap=False, fill=None):
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=self.ncols)
        c = ws.cell(self.r, 1, t(value) if isinstance(value, dict) else value)
        c.font = self.f(size, bold, color)
        c.alignment = Alignment(horizontal=self.align, vertical='center', wrap_text=wrap)
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)
        if height:
            ws.row_dimensions[self.r].height = height
        self.r += 1

    def label(self, value):
        self.text(value, size=10, bold=True, color=AMBER_D, height=20)

    def header(self, cols):
        for i, h in enumerate(cols, 1):
            c = self.ws.cell(self.r, i, t(h))
            c.font = self.f(9, True, 'FFFFFF')
            c.fill = PatternFill('solid', fgColor=DARK)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = BORDER
        self.ws.row_dimensions[self.r].height = 30
        self.r += 1

    def rows(self, n, height=30, pre=None, numbered=False, formulas=None, fills=None):
        """n empty rows; pre: first-column values; formulas: {col: template with {r}}."""
        first = self.r
        for k in range(n):
            for i in range(1, self.ncols + 1):
                c = self.ws.cell(self.r, i)
                c.border = BORDER
                c.alignment = Alignment(horizontal=self.align, vertical='top', wrap_text=True)
                c.font = self.f(10)
                if fills and i in fills:
                    c.fill = PatternFill('solid', fgColor=fills[i])
            if pre and k < len(pre):
                c = self.ws.cell(self.r, 1, t(pre[k]))
                c.font = self.f(10, True)
                c.fill = PatternFill('solid', fgColor='FAF8F4')
            elif numbered:
                c = self.ws.cell(self.r, 1, k + 1)
                c.font = self.f(10, True, AMBER_D)
                c.alignment = Alignment(horizontal='center', vertical='top')
            for col, tpl in (formulas or {}).items():
                c = self.ws.cell(self.r, col, tpl.format(r=self.r))
                c.font = self.f(10, True)
                c.fill = PatternFill('solid', fgColor=CREAM)
                c.alignment = Alignment(horizontal='center', vertical='top')
            self.ws.row_dimensions[self.r].height = height
            self.r += 1
        return first, self.r - 1

    def table(self, cols, n, **kw):
        self.header(cols)
        return self.rows(n, **kw)

    def field(self, label, height=60):
        """A labelled writing box spanning the sheet."""
        self.label(label)
        self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=self.ncols)
        c = self.ws.cell(self.r, 1)
        c.alignment = Alignment(horizontal=self.align, vertical='top', wrap_text=True)
        c.font = self.f(10)
        for i in range(1, self.ncols + 1):
            self.ws.cell(self.r, i).border = BORDER
            self.ws.cell(self.r, i).fill = PatternFill('solid', fgColor='FFFFFF')
        self.ws.row_dimensions[self.r].height = height
        self.r += 2

    def choice(self, rng, options):
        dv = DataValidation(type='list', formula1='"' + ','.join(t(o) for o in options) + '"', allow_blank=True)
        self.ws.add_data_validation(dv)
        dv.add(rng)

    def whole(self, rng, lo, hi):
        dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True)
        dv.error = t(bi(f'Enter a whole number from {lo} to {hi}.', f'أدخل عدداً صحيحاً من {lo} إلى {hi}.'))
        self.ws.add_data_validation(dv)
        dv.add(rng)

    def gap(self, n=1):
        self.r += n


def col(i):
    return get_column_letter(i)


YES, NOT_YET = bi('Yes', 'نعم'), bi('Not yet', 'ليس بعد')
OWNER = bi('Owner', 'المسؤول')


def build(lang):
    global L
    L = lang
    stages = C['stages']
    S = {s['key']: s for s in stages}
    wb = Workbook()
    wb.remove(wb.active)

    # Start
    sh = Sheet(wb, None, bi('Start here', 'ابدأ من هنا'), C['meta']['subtitle'], [26, 70])
    sh.text(t(C['meta']['author']), size=11, bold=True)
    sh.text('www.fawzooz.ai', size=10, color=AMBER_D)
    sh.gap()
    sh.label(bi('How to use it', 'كيف تستخدمه'))
    for line in [
        bi('Fill in the white cells. Cream cells are calculated for you.', 'املأ الخلايا البيضاء؛ أما الخلايا الكريمية فتُحسب تلقائياً.'),
        bi('Open with the canvas and Ignite together at kick-off; revisit the canvas at every milestone.', 'ابدأ باللوحة وبالإشعال معاً عند الانطلاق، وعُد إلى اللوحة عند كل محطة.'),
        bi('Close each stage on the Gates sheet before moving on. If a criterion is not met, loop back.', 'أغلق كل مرحلة في ورقة البوابات قبل الانتقال؛ فإن لم يتحقق معيار، فعُد إلى الوراء.'),
        bi('The printable edition (PDF) has every sheet, including those drawn by hand: the mind map, the storyboard, the sketch.', 'تضم النسخة المطبوعة (PDF) كل الأوراق، ومنها ما يُرسم باليد: الخريطة الذهنية، واللوحة القصصية، والرسم.'),
    ]:
        sh.text('•  ' + t(line), wrap=True, height=30)
    sh.gap()
    sh.label(bi('The five stages', 'المراحل الخمس'))
    sh.header([bi('Stage', 'المرحلة'), bi('Exit criteria', 'معايير الخروج')])
    for s in stages:
        sh.ws.cell(sh.r, 1, f"{num(s['n'])} · {t(s['name'])}").font = sh.f(10, True)
        c = sh.ws.cell(sh.r, 2, t(s['exit']))
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal=sh.align)
        c.font = sh.f(10)
        for i in (1, 2):
            sh.ws.cell(sh.r, i).border = BORDER
        sh.ws.row_dimensions[sh.r].height = 34
        sh.r += 1
    sh.gap()
    sh.text(t(bi('Licence: CC BY-SA 4.0. You may copy, adapt and share these templates, including commercially, with credit and under the same licence.',
                 'الترخيص: CC BY-SA 4.0. يحق لك نسخ هذه القوالب وتعديلها ومشاركتها، ولو لأغراض تجارية، مع نسبتها إلى مصدرها وبالرخصة نفسها.')), size=9, color='8E8B85', wrap=True, height=30)

    # 0.2 Canvas
    sh = Sheet(wb, '0.2', bi('Canvas', 'اللوحة'), bi('The single source of truth: one row per stage, three boxes each.', 'المرجع الموحّد: صف لكل مرحلة، وثلاث خانات لكلٍّ منها.'), [16, 38, 38, 38])
    sh.ws.cell(sh.r, 1, t(bi('Project', 'المشروع'))).font = sh.f(10, True)
    sh.ws.merge_cells(start_row=sh.r, start_column=2, end_row=sh.r, end_column=4)
    sh.ws.cell(sh.r, 2).border = BORDER
    sh.r += 2
    for i, s in enumerate(stages):
        sh.header([f"{num(s['n'])} · {t(s['name'])}"] + [f[L] for f in C['canvas'][i]])
        sh.rows(1, height=110, pre=[s['tagline']])

    # 1.1 Five Whys
    sh = Sheet(wb, '1.1', bi('Five Whys', 'اللماذات الخمس'), bi('Start from a fact you have seen; ask “why?” of each answer.', 'ابدأ من حقيقة رأيتها، واسأل «لماذا؟» عن كل جواب.'), [16, 90])
    sh.field(bi('The problem we observed', 'المشكلة التي لاحظناها'), 50)
    sh.table([bi('Why?', 'لماذا؟'), bi('Because…', 'لأن…')], 5, height=36, pre=[f"{t(bi('Why?', 'لماذا؟'))} {num(i)}" for i in range(1, 6)])
    sh.gap()
    sh.field(bi('Root cause, in one sentence', 'السبب الجذري، في جملة واحدة'), 40)

    # 1.2 Inspiration
    sh = Sheet(wb, '1.2', bi('Inspiration', 'الإلهام'), bi('Reports, competitors, customers and unrelated fields: what you saw and what it suggests.', 'التقارير والمنافسون والعملاء والمجالات البعيدة: ما رأيت وما يوحي به.'), [26, 44, 44])
    sh.table([bi('Source', 'المصدر'), bi('What we saw', 'ما رأيناه'), bi('What it suggests', 'ما يوحي به')], 12, height=40,
             pre=[bi('Reports and research', 'التقارير والأبحاث'), bi('Competitors', 'المنافسون'), bi('Customers', 'العملاء'), bi('Customers', 'العملاء'), bi('An unrelated field', 'مجال بعيد'), bi('An unrelated field', 'مجال بعيد')])
    sh.gap()
    sh.field(bi('The pattern across sources', 'النمط المشترك بين المصادر'), 50)

    # 1.3 Idea bank
    sh = Sheet(wb, '1.3', bi('Idea bank', 'رصيد الأفكار'), bi('Defer judgment, build on others’ ideas, go for quantity. Vote only at the end.', 'أرجئ الحكم، وابنِ على أفكار الآخرين، واطلب الكمّ. ولا تصوّت إلا في النهاية.'), [6, 70, 18, 10])
    sh.field(bi('How might we…?', 'كيف يمكننا…؟'), 30)
    a, b = sh.table(['#', bi('Idea', 'الفكرة'), bi('From (brainstorm, provocation…)', 'المصدر (عصف، استفزاز…)'), bi('Votes', 'الأصوات')], 40, height=22, numbered=True)
    sh.choice(f'C{a}:C{b}', [bi('Brainstorm', 'العصف الذهني'), bi('Provocation', 'الاستفزاز الذهني'), bi('Inspiration', 'الإلهام'), bi('Five Whys', 'اللماذات الخمس')])

    # 1.4 Prioritize
    sh = Sheet(wb, '1.4', bi('Prioritize', 'الأولويات'), bi('Score 1 (low) to 5 (high). The total and the rank are calculated; keep the top one to three.', 'قيّم من ١ (منخفض) إلى ٥ (مرتفع). يُحسب المجموع والترتيب تلقائياً؛ واحتفظ بالأعلى من واحد إلى ثلاثة.'), [46, 11, 11, 11, 11, 11, 9])
    a, b = sh.table([bi('Problem or opportunity', 'المشكلة أو الفرصة'), bi('Impact', 'الأثر'), bi('Urgency', 'الإلحاح'), bi('Evidence', 'الدليل'), bi('Fit', 'المواءمة'), bi('Total', 'المجموع'), bi('Rank', 'الترتيب')], 10, height=28,
                    formulas={6: '=IF(COUNT(B{r}:E{r})=0,"",SUM(B{r}:E{r}))', 7: ''})
    for r in range(a, b + 1):
        sh.ws.cell(r, 7).value = f'=IF(F{r}="","",COUNTIF($F${a}:$F${b},">"&F{r})+1)'
    sh.whole(f'B{a}:E{b}', 1, 5)
    sh.ws.conditional_formatting.add(f'G{a}:G{b}', CellIsRule(operator='lessThanOrEqual', formula=['3'], fill=PatternFill('solid', fgColor=AMBER), font=Font(bold=True, color=DARK)))

    # 2.1 Concept and Lean Canvas
    sh = Sheet(wb, '2.1', bi('Concept', 'المفهوم'), bi('Purpose, audience and features; then the value in one sentence, and a thirty-minute Lean Canvas.', 'الغاية والجمهور والخصائص، ثم القيمة في جملة واحدة، ولوحة رشيقة في ثلاثين دقيقة.'), [30, 76])
    for lab, h in [(bi('Purpose: what it is for', 'الغاية: لأيّ شيء هو'), 40), (bi('Audience: who it is for', 'الجمهور: لمن هو'), 40), (bi('Features: what it does', 'الخصائص: ما الذي يفعله'), 60),
                   (bi('Concept statement: For … who …, … is a … that … . Unlike …, it … .', 'بيان المفهوم: لـ… الذين…، فإن… هو… الذي… . وبخلاف…، فإنه… .'), 60), (bi('Value proposition, in one sentence', 'عرض القيمة، في جملة واحدة'), 36)]:
        sh.field(lab, h)
    sh.label(bi('Lean Canvas (after Maurya, 2012)', 'اللوحة الرشيقة (عن موريا، ٢٠١٢)'))
    a, b = sh.table([bi('Box', 'الخانة'), bi('Our answer', 'جوابنا')], 9, height=44, pre=[
        bi('1 Customer segments', '١ شرائح العملاء'), bi('2 Problem', '٢ المشكلة'), bi('3 Unique value proposition', '٣ عرض القيمة الفريد'), bi('4 Solution', '٤ الحل'), bi('5 Channels', '٥ القنوات'),
        bi('6 Revenue streams', '٦ مصادر الإيراد'), bi('7 Cost structure', '٧ هيكل التكاليف'), bi('8 Key metrics', '٨ المقاييس الرئيسية'), bi('9 Unfair advantage', '٩ الميزة التي لا تُنسخ')])

    # 2.2 SCAMPER
    sh = Sheet(wb, '2.2', bi('SCAMPER', 'سكامبر'), bi('At least one answer per prompt; star the two or three that change the value most.', 'جواب واحد على الأقل لكل محفّز، وضع نجمة على الاثنين أو الثلاثة الأكثر تغييراً للقيمة.'), [24, 70, 8])
    sh.field(bi('The idea we are working on', 'الفكرة التي نعمل عليها'), 30)
    a, b = sh.table([bi('Prompt', 'المحفّز'), bi('Our answers', 'أجوبتنا'), '★'], 7, height=44, pre=[
        bi('S · Substitute', 'S · استبدل'), bi('C · Combine', 'C · ادمج'), bi('A · Adapt', 'A · كيّف'), bi('M · Modify', 'M · عدّل'), bi('P · Put to another use', 'P · استخدم لغرض آخر'), bi('E · Eliminate', 'E · احذف'), bi('R · Reverse', 'R · اعكس')])
    sh.choice(f'C{a}:C{b}', ['★'])

    # 3.1 Assumptions
    sh = Sheet(wb, '3.1', bi('Assumptions', 'الافتراضات'), bi('Importance and evidence from 1 to 5. Important assumptions with weak evidence come first.', 'الأهمية والدليل من ١ إلى ٥. الافتراضات المهمّة ذات الدليل الضعيف تأتي أولاً.'), [5, 56, 16, 12, 12, 12, 14])
    a, b = sh.table(['#', bi('We believe that…', 'نعتقد أن…'), bi('Type', 'النوع'), bi('Importance 1–5', 'الأهمية ١–٥'), bi('Evidence 1–5', 'الدليل ١–٥'), bi('Risk score', 'درجة الخطر'), bi('Quadrant', 'الربع')], 15, height=30, numbered=True,
                    formulas={6: '=IF(OR(D{r}="",E{r}=""),"",D{r}*(6-E{r}))', 7: '=IF(F{r}="","",IF(D{r}>=3,IF(E{r}<=3,"' + t(bi('Test first', 'اختبر أولاً')) + '","' + t(bi('Watch', 'راقب')) + '"),IF(E{r}<=3,"' + t(bi('Later', 'لاحقاً')) + '","' + t(bi('Park it', 'تجاهل الآن')) + '")))'})
    sh.choice(f'C{a}:C{b}', [bi('Desirable', 'مرغوب'), bi('Feasible', 'ممكن'), bi('Viable', 'مُجدٍ')])
    sh.whole(f'D{a}:E{b}', 1, 5)
    sh.ws.conditional_formatting.add(f'G{a}:G{b}', FormulaRule(formula=[f'G{a}="{t(bi("Test first", "اختبر أولاً"))}"'], fill=PatternFill('solid', fgColor=AMBER), font=Font(bold=True)))

    # 3.2 Tests
    sh = Sheet(wb, '3.2', bi('Tests', 'الاختبارات'), bi('One row per test. Write “We are right if…” before the test runs.', 'صفّ لكل اختبار. واكتب «نكون على حق إذا…» قبل إجراء الاختبار.'), [5, 10, 34, 16, 10, 22, 28, 28, 14])
    a, b = sh.table(['#', bi('Assumption #', 'رقم الافتراض'), bi('We believe that…', 'نعتقد أن…'), bi('Method', 'الطريقة'), bi('Users', 'المستخدمون'), bi('We measure', 'نقيس'), bi('We are right if…', 'نكون على حق إذا…'), bi('What happened', 'ما حدث'), bi('Verdict', 'الحكم')], 12, height=36, numbered=True)
    sh.choice(f'D{a}:D{b}', [bi('Interviews', 'مقابلات'), bi('Analytics', 'تحليلات'), bi('A/B test', 'اختبار مقارنة'), bi('Usability test', 'اختبار سهولة الاستخدام')])
    sh.choice(f'I{a}:I{b}', [bi('Validated', 'مُتحقَّق منه'), bi('Invalidated', 'مدحوض')])
    sh.whole(f'E{a}:E{b}', 0, 10000)

    # 3.3 Feedback
    sh = Sheet(wb, '3.3', bi('Feedback log', 'سجل الملاحظات'), bi('Five to ten users per cycle. What they did comes before what they said.', 'من خمسة إلى عشرة مستخدمين في كل دورة. ما فعلوه قبل ما قالوه.'), [8, 16, 34, 34, 24, 14])
    a, b = sh.table([bi('Cycle', 'الدورة'), bi('User', 'المستخدم'), bi('What they did', 'ما فعلوه'), bi('What they said', 'ما قالوه'), bi('Surprise', 'المفاجأة'), bi('Assumption ±', 'الافتراض ±')], 30, height=30)

    # 3.5 Cycles
    sh = Sheet(wb, '3.5', bi('Build–Measure–Learn', 'ابنِ ثم قِس ثم تعلّم'), bi('Every cycle ends in a decision on evidence: persevere, or pivot.', 'تنتهي كل دورة بقرار مبنيّ على الدليل: المضيّ أو التحوّل.'), [8, 30, 30, 34, 14])
    a, b = sh.table([bi('Cycle', 'الدورة'), bi('What we built', 'ما بنيناه'), bi('What we measured', 'ما قسناه'), bi('What we learned', 'ما تعلّمناه'), bi('Decision', 'القرار')], 8, height=44, numbered=True)
    sh.choice(f'E{a}:E{b}', [bi('Persevere', 'نمضي'), bi('Pivot', 'نتحوّل')])
    sh.gap()
    sh.label(bi('Pivot record', 'سجل التحوّلات'))
    sh.table([bi('From', 'من'), bi('To', 'إلى'), bi('The evidence that made us turn', 'الدليل الذي جعلنا نتحوّل'), '', ''], 5, height=36)

    # 4.1 SWOT
    sh = Sheet(wb, '4.1', bi('SWOT', 'التحليل الرباعي'), bi('Tie the solution to the mission first; then pair the quadrants into roadmap items and risks.', 'اربط الحل بالرسالة أولاً، ثم زاوج بين الأرباع لتخرج ببنود لخارطة الطريق وبالمخاطر.'), [50, 50])
    sh.field(bi('The line of our mission or strategy it serves, and how', 'سطر الرسالة أو الاستراتيجية الذي يخدمه، وكيف'), 44)
    sh.header([bi('Strengths (inside)', 'نقاط القوة (داخلية)'), bi('Weaknesses (inside)', 'نقاط الضعف (داخلية)')])
    sh.rows(1, height=120)
    sh.header([bi('Opportunities (outside)', 'الفرص (خارجية)'), bi('Threats (outside)', 'التهديدات (خارجية)')])
    sh.rows(1, height=120)
    sh.header([bi('Strength × opportunity → roadmap item', 'قوة × فرصة ← بند في خارطة الطريق'), bi('Weakness × threat → risk', 'ضعف × تهديد ← خطر')])
    sh.rows(1, height=80)

    # 4.3 OKRs
    sh = Sheet(wb, '4.3', bi('OKRs', 'الأهداف والنتائج'), bi('Each key result needs a baseline and a target. Enter the current value to see progress; Polish tracks these.', 'لكل نتيجة رئيسية خط أساس وقيمة مستهدفة. أدخل القيمة الحالية لترى التقدّم؛ والصقل يتابعها.'), [30, 40, 11, 11, 11, 11, 14, 12])
    a, b = sh.table([bi('Objective', 'الهدف'), bi('Key result', 'النتيجة الرئيسية'), bi('Baseline', 'خط الأساس'), bi('Target', 'المستهدف'), bi('Current', 'الحالي'), bi('Progress', 'التقدّم'), OWNER, bi('By', 'بحلول')], 12, height=30,
                    formulas={6: '=IF(OR(C{r}="",D{r}="",E{r}="",D{r}=C{r}),"",(E{r}-C{r})/(D{r}-C{r}))'})
    for r in range(a, b + 1):
        sh.ws.cell(r, 6).number_format = '0%'
    for k in range(4):
        sh.ws.merge_cells(start_row=a + 3 * k, start_column=1, end_row=a + 3 * k + 2, end_column=1)

    # 4.4 Roadmap
    months = [f"{t(bi('Month', 'الشهر'))} {num(i)}" for i in range(1, 7)]
    sh = Sheet(wb, '4.4', bi('Roadmap', 'خارطة الطريق'), bi('Three to six months by workstream; then the first sprints as user stories.', 'من ثلاثة إلى ستة أشهر بحسب مسارات العمل، ثم دورات العمل الأولى قصصَ مستخدم.'), [22, 18, 18, 18, 18, 18, 18])
    sh.table([bi('Workstream', 'المسار')] + months, 5, height=50, pre=[bi('Product', 'المنتج'), bi('People and skills', 'الأفراد والمهارات'), bi('Users and market', 'المستخدمون والسوق'), bi('Operations', 'التشغيل'), bi('Milestones', 'المحطات')])
    sh.gap()
    sh.label(bi('Sprints and user stories', 'دورات العمل وقصص المستخدم'))
    a, b = sh.table([bi('Sprint', 'الدورة'), bi('As a…, I want…, so that…', 'بصفتي…، أريد…، كي…'), '', '', bi('Done when', 'تُنجَز حين'), '', bi('Status', 'الحالة')], 20, height=28)
    for r in range(a, b + 1):
        sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        sh.ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    sh.choice(f'G{a}:G{b}', [bi('To do', 'لم يبدأ'), bi('Doing', 'جارٍ'), bi('Done', 'أُنجز')])

    # 4.5 Resources and risks
    sh = Sheet(wb, '4.5', bi('Resources and risks', 'الموارد والمخاطر'), bi('Likelihood and impact from 1 to 3; the score is calculated. Six and above: act now.', 'الاحتمال والأثر من ١ إلى ٣، وتُحسب الدرجة تلقائياً. ستة فما فوق: تصرّف الآن.'), [40, 10, 10, 10, 34, 16])
    a, b = sh.table([bi('Resource: people, tools, space', 'المورد: أفراد، وأدوات، ومكان'), bi('How much', 'الكمية'), bi('Cost', 'الكلفة'), '', '', ''], 8, height=24)
    sh.ws.cell(sh.r, 1, t(bi('Total budget', 'إجمالي الميزانية'))).font = sh.f(10, True)
    c = sh.ws.cell(sh.r, 3, f'=SUM(C{a}:C{b})')
    c.font = sh.f(10, True)
    c.fill = PatternFill('solid', fgColor=CREAM)
    sh.r += 2
    a, b = sh.table([bi('Risk', 'الخطر'), bi('Likelihood 1–3', 'الاحتمال ١–٣'), bi('Impact 1–3', 'الأثر ١–٣'), bi('Score', 'الدرجة'), bi('Response', 'الاستجابة'), OWNER], 12, height=30,
                    formulas={4: '=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})'})
    sh.whole(f'B{a}:C{b}', 1, 3)
    for op, f, fill in [('greaterThanOrEqual', ['6'], 'F3C3B3'), ('between', ['3', '4'], 'FBE7BF'), ('between', ['1', '2'], 'E9EFE3')]:
        sh.ws.conditional_formatting.add(f'D{a}:D{b}', CellIsRule(operator=op, formula=f, fill=PatternFill('solid', fgColor=fill)))

    # 5.1 Launch checklist
    sh = Sheet(wb, '5.1', bi('Launch checklist', 'قائمة الإطلاق'), bi('Nothing ships until every line is Yes or consciously waived.', 'لا يُطلَق شيء حتى تكون كل خانة «نعم» أو يُتنازل عنها عن وعي.'), [70, 14, 20])
    sh.field(bi('Kind of launch: full, or phased (who gets it, from when, and go on if…)', 'نوع الإطلاق: كامل، أو تدريجي (مَن يحصل عليه، ومن متى، ونستمر إذا…)'), 50)
    items = [
        bi('The key results from Sculpt have a baseline.', 'للنتائج الرئيسية من النحت خط أساس.'), bi('Tracking is live, and someone owns the dashboard.', 'المتابعة تعمل، وللوحة المؤشرات مسؤول.'),
        bi('Users have a channel for feedback, and someone reads it.', 'للمستخدمين قناة للملاحظات، وثمة من يقرؤها.'), bi('Support knows what is coming and how to answer.', 'يعرف فريق الدعم ما القادم وكيف يجيب.'),
        bi('There is a way back if something breaks.', 'ثمة طريق للرجوع إن تعطّل شيء.'), bi('Someone is on call for the first days.', 'ثمة مناوب في الأيام الأولى.'),
        bi('Legal, privacy and security have signed off.', 'وافقت الجهات القانونية والخصوصية والأمن.'), bi('Stakeholders and users have been told.', 'أُبلغ أصحاب المصلحة والمستخدمون.'),
        bi('The retrospective is booked, two to four weeks out.', 'حُجز موعد المراجعة اللاحقة بعد أسبوعين إلى أربعة.'), bi('The canvas is up to date.', 'اللوحة محدّثة.'),
        '', '', '']
    a, b = sh.table([bi('Item', 'البند'), bi('Done', 'أُنجز'), OWNER], len(items), height=24, pre=items)
    sh.choice(f'B{a}:B{b}', [YES, bi('Waived', 'تُنوزل عنه'), NOT_YET])
    sh.ws.cell(sh.r, 1, t(bi('Ready', 'الجاهزية'))).font = sh.f(10, True)
    c = sh.ws.cell(sh.r, 2, f'=COUNTIF(B{a}:B{b},"{t(YES)}")+COUNTIF(B{a}:B{b},"{t(bi("Waived", "تُنوزل عنه"))}")&" / "&COUNTA(A{a}:A{b})')
    c.font = sh.f(10, True)
    c.fill = PatternFill('solid', fgColor=CREAM)

    # 5.2 Impact tracker
    sh = Sheet(wb, '5.2', bi('Impact tracker', 'متابعة الأثر'), bi('Track the key results and KPIs from Sculpt from day one. Read the trend, not a single number.', 'تابع النتائج الرئيسية ومؤشرات الأداء من النحت منذ اليوم الأول، واقرأ الاتجاه لا رقماً منفرداً.'), [40, 11, 11, 11, 11, 11, 12, 12])
    a, b = sh.table([bi('Key result or KPI', 'النتيجة الرئيسية أو المؤشر'), bi('Baseline', 'الأساس'), bi('Target', 'المستهدف'), bi('Week 1', 'الأسبوع ١'), bi('Week 2', 'الأسبوع ٢'), bi('Week 4', 'الأسبوع ٤'), bi('Progress', 'التقدّم'), bi('Trend', 'الاتجاه')], 10, height=26,
                    formulas={7: '=IF(OR(B{r}="",C{r}="",C{r}=B{r},COUNT(D{r}:F{r})=0),"",(IF(F{r}<>"",F{r},IF(E{r}<>"",E{r},D{r}))-B{r})/(C{r}-B{r}))',
                              8: '=IF(COUNT(D{r}:F{r})<2,"",IF(IF(F{r}<>"",F{r},E{r})>D{r},"▲",IF(IF(F{r}<>"",F{r},E{r})<D{r},"▼","▬")))'})
    for r in range(a, b + 1):
        sh.ws.cell(r, 7).number_format = '0%'
    sh.gap()
    sh.field(bi('What the numbers say', 'ما تقوله الأرقام'), 60)

    # 5.3 Retrospective
    sh = Sheet(wb, '5.3', bi('Retrospective', 'المراجعة اللاحقة'), bi('Two to four weeks after launch, with the whole team.', 'بعد أسبوعين إلى أربعة أسابيع من الإطلاق، مع الفريق كله.'), [36, 36, 36])
    sh.header([bi('What worked', 'ما الذي نجح'), bi('What did not', 'ما الذي لم ينجح'), bi('What we learned', 'ماذا تعلّمنا')])
    sh.rows(1, height=200)
    sh.gap()
    sh.label(bi('Actions', 'الإجراءات'))
    sh.table([bi('Action', 'الإجراء'), OWNER, bi('By when', 'الموعد')], 8, height=26)

    # 5.4 Decision
    sh = Sheet(wb, '5.4', bi('Decision', 'القرار'), bi('Decide on the data, then hand the next Ignite what this cycle found.', 'قرّر بناءً على البيانات، ثم سلّم الإشعال التالي ما وجدته هذه الدورة.'), [30, 76])
    sh.label(bi('The decision', 'القرار'))
    sh.ws.cell(sh.r, 1).border = BORDER
    sh.choice(f'A{sh.r}', [bi('Iterate', 'التكرار'), bi('Scale', 'التوسّع'), bi('Sunset', 'الإيقاف')])
    sh.ws.cell(sh.r, 2, t(bi('Choose: iterate, scale or sunset', 'اختر: التكرار أو التوسّع أو الإيقاف'))).font = sh.f(9, color='8E8B85')
    sh.r += 2
    for lab, h in [(bi('The evidence behind the decision', 'الدليل وراء القرار'), 70), (bi('What scaling, iterating or closing will need', 'ما يحتاجه التوسّع أو التكرار أو الإغلاق'), 50),
                   (bi('Seeds for the next Ignite: problems this cycle uncovered', 'بذور الإشعال التالي: مشكلات كشفتها هذه الدورة'), 50), (bi('Ideas we parked', 'أفكار أرجأناها'), 50), (bi('What we will do differently', 'ما سنفعله على نحو مختلف'), 50)]:
        sh.field(lab, h)

    # Gates
    sh = Sheet(wb, '✓', bi('Gates', 'البوابات'), bi('Close each stage only when every criterion is Yes.', 'لا تُغلق المرحلة إلا حين تكون معاييرها كلها «نعم».'), [70, 12, 30])
    for s in stages:
        sh.label(f"{t(bi('Stage', 'المرحلة'))} {num(s['n'])} · {t(s['name'])}")
        a, b = sh.table([bi('Criterion', 'المعيار'), bi('Met?', 'تحقق؟'), bi('Evidence', 'الدليل')], len(s['gate']), height=30, pre=s['gate'])
        sh.choice(f'B{a}:B{b}', [YES, NOT_YET])
        sh.ws.cell(sh.r, 1, t(bi('Gate', 'البوابة'))).font = sh.f(10, True)
        open_ = t(bi('Open', 'مفتوحة'))
        closed = t(bi('Closed: move on', 'مُغلقة: انتقل'))
        c = sh.ws.cell(sh.r, 2, f'=IF(COUNTIF(B{a}:B{b},"{t(YES)}")={b - a + 1},"{closed}","{open_} "&COUNTIF(B{a}:B{b},"{t(YES)}")&"/{b - a + 1}")')
        c.font = sh.f(10, True)
        c.fill = PatternFill('solid', fgColor=CREAM)
        sh.ws.merge_cells(start_row=sh.r, start_column=2, end_row=sh.r, end_column=3)
        sh.r += 2

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.title = t(C['meta']['title'])
    wb.properties.creator = t(C['meta']['author'])
    path = os.path.join(OUT, f'innovaforge-toolkit-{lang}.xlsx')
    wb.save(path)
    print(os.path.relpath(path, DIR), len(wb.worksheets), 'sheets')


for lang in ('en', 'ar'):
    build(lang)
